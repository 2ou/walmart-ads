"""Check the 'Trade Item Configurations' sheet for empty WFS-required fields."""
from openpyxl import load_workbook

DST = r"D:\Tools\Code\walmart-ads\outputs\PNK366085-Blouses-Tops-upload-optimized-20260916.xlsx"
LOG = r"D:\WorkBuddyData\walmart-ads\listing_dump\trade_item_check.txt"

wb = load_workbook(DST, read_only=True)
ws = wb["Trade Item Configurations"]
rows = [list(r) for r in ws.iter_rows(values_only=True)]
wb.close()


def s(v):
    return "" if v is None else str(v).strip()


id_idx = next(i for i, r in enumerate(rows) if len(r) > 3 and s(r[3]) == "sku")
ids = [s(x) for x in rows[id_idx]]
labels = [s(x) for x in rows[id_idx - 1]]
spec = [s(x) for x in rows[id_idx + 1]]
first = rows[id_idx + 2]

L = ["attrId row idx=%d" % id_idx, ""]
for i, a in enumerate(ids):
    if not a:
        continue
    L.append("%-28s | %-26s | val=%-18s | %s" % (
        a, labels[i][:26], s(first[i])[:18], spec[i][:70]))

L.append("")
L.append("=== empty columns ===")
for i, a in enumerate(ids):
    if a and not s(first[i]):
        L.append("  %-28s %s" % (a, labels[i][:40]))

with open(LOG, "w", encoding="utf-8") as f:
    f.write("\n".join(L))
print("WROTE", LOG)
