"""Inspect the Walmart MP_WFS_ITEM upload template structure before patching.

Outputs a report to D:\\WorkBuddyData\\walmart-ads\\listing_dump\\template_inspect.txt
"""
import os
from openpyxl import load_workbook

SRC = r"C:\Users\Administrator\Desktop\PNK366085-Blouses-Tops-upload.xlsx"
OUT = r"D:\WorkBuddyData\walmart-ads\listing_dump\template_inspect.txt"

FILLABLE_IDS = [
    "clothing_shoulder_style", "collarType", "sleeveStyle", "clothing_back_style",
    "manufacturer", "manufacturerPartNumber", "modelNumber", "productLine", "theme",
    "netContentStatement", "ib_retail_packaging", "shipsInOriginalPackaging", "designer",
    "msrp", "minimumSellerAllowedPrice", "maximumSellerAllowedPrice", "repricerStrategy",
    "recycledMaterial", "percentageOfRecycledMaterial", "warrantyText", "warrantyURL",
    "SkuUpdate", "externalProductIdType", "externalProductId",
    "variantAttributeNames", "productName", "colorCategory", "isPrimaryVariant",
    "clothingSize", "clothingSizeGroup", "keyFeatures", "productSecondaryImageURL",
]


def s(v):
    return "" if v is None else str(v).strip()


wb = load_workbook(SRC, read_only=True)
L = []
L.append("SHEETS: %s" % wb.sheetnames)

ws = wb["Product Content And Site Exp"]
rows = [list(r) for r in ws.iter_rows(values_only=True)]
wb.close()
L.append("dims: %d rows x %d cols" % (len(rows), len(rows[0]) if rows else 0))

# locate the attribute-id row: exactly lowercase "sku" in col 4
id_idx = None
for i, r in enumerate(rows):
    if len(r) > 3 and s(r[3]) == "sku":
        id_idx = i
        break
L.append("attribute-id row index (0-based) = %s" % id_idx)
for k in range(max(0, id_idx - 3), min(len(rows), id_idx + 3)):
    L.append("  row[%d] first12: %s" % (k, [s(x) for x in rows[k][:12]]))

ids = [s(x) for x in rows[id_idx]]
labels = [s(x) for x in rows[id_idx - 1]]
help_row = rows[id_idx + 1] if id_idx + 1 < len(rows) else []
spec = [s(x) for x in help_row]

first = rows[id_idx + 2]

L.append("")
L.append("=== ALL COLUMNS (idx | label | attrId | value_in_row1 | spec) ===")
for i in range(len(ids)):
    if not ids[i] and not labels[i]:
        continue
    L.append("%3d | %-46s | %-44s | %-40s | %s" % (
        i, labels[i][:46], ids[i][:44], s(first[i])[:40], spec[i][:60] if i < len(spec) else ""))

L.append("")
L.append("=== FILLABLE ID LOOKUP (occurrence 1) ===")
for fid in FILLABLE_IDS:
    hits = [i for i, a in enumerate(ids) if a == fid]
    if not hits:
        L.append("  %-40s -> NOT FOUND" % fid)
    else:
        vals = [s(rows[id_idx + 2 + r][hits[0]]) for r in range(min(3, len(rows) - id_idx - 2))]
        L.append("  %-40s -> cols %s | label=%s | sample=%s" % (
            fid, hits, labels[hits[0]][:40], vals))

L.append("")
L.append("=== DATA ROWS (%d) ===" % (len(rows) - id_idx - 2))
def find_col(name):
    for i, a in enumerate(ids):
        if a.lower() == name.lower():
            return i
    return None


sku_c = find_col("sku")
for r in range(id_idx + 2, len(rows)):
    rr = rows[r]
    if len(rr) <= sku_c or not s(rr[sku_c]):
        continue
    col_c = find_col("color")
    sz_c = find_col("clothingSize")
    L.append("  row%-3d sku=%-38s color=%-14s size=%-4s" % (
        r + 1, s(rr[sku_c]), s(rr[col_c]) if col_c is not None else "?", s(rr[sz_c]) if sz_c is not None else "?"))

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(L))
print("WROTE", OUT)
