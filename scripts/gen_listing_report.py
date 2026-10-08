"""Generate the PNK366085 Walmart listing optimization report (offline single-file HTML).

Reads:
  - PNK366085-Blouses-Tops-upload.xlsx           (own listing)
  - ReverseASIN-US-B0F93MLVNZ(2000)-20260916.xlsx (competitor keyword universe)
  - US-Alexa(5)-20260916.xlsx                     (competitor AI Q&A + compare set)
Writes:
  outputs/PNK366085_Listing优化方案_2026-09-16.html
"""
import json
import os
import re
from openpyxl import load_workbook

DESK = r"C:\Users\Administrator\Desktop"
SELF_X = os.path.join(DESK, "PNK366085-Blouses-Tops-upload.xlsx")
REV_X = os.path.join(DESK, "ReverseASIN-US-B0F93MLVNZ(2000)-20260916.xlsx")
ALX_X = os.path.join(DESK, "US-Alexa(5)-20260916.xlsx")
ECHARTS = r"D:\Tools\Code\GoodJob\node_modules\echarts\dist\echarts.min.js"
OUT = r"D:\Tools\Code\walmart-ads\outputs\PNK366085_Listing优化方案_2026-09-16.html"


def load(path, sheet=None):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def s(v):
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------- own listing
self_rows = load(SELF_X, "Product Content And Site Exp")
id_idx = next(i for i, r in enumerate(self_rows) if len(r) > 3 and s(r[3]) == "sku")
ids = [s(x) for x in self_rows[id_idx]]
labels = [s(x) for x in self_rows[id_idx - 1]]
data = [r for r in self_rows[id_idx + 2:] if len(r) > 3 and s(r[3])]


def col(name, n=1):
    cnt = 0
    for i, a in enumerate(ids):
        if a == name:
            cnt += 1
            if cnt == n:
                return i
    return None


first = data[0]
C = {k: col(k) for k in [
    "sku", "productName", "brand", "price", "shortDescription", "mainImageUrl",
    "color", "clothingSize", "clothingSizeGroup", "colorCategory", "materialName",
    "materialPercentage", "variantAttributeNames", "isPrimaryVariant"]}
KF = [col("keyFeatures", n) for n in range(1, 7)]
SEC = [col("productSecondaryImageURL", n) for n in range(1, 9)]

TITLE = s(first[C["productName"]])
BULLETS = [s(first[i]) for i in KF if i is not None]
DESC = s(first[C["shortDescription"]])
COLORS = sorted({s(r[C["color"]]) for r in data})
SIZES = sorted({s(r[C["clothingSize"]]) for r in data})
PRICE = s(first[C["price"]])

# own searchable text (title + bullets + description, tags stripped)
OWN_TEXT = " ".join([TITLE] + BULLETS + [re.sub(r"<[^>]+>", " ", DESC)]).lower()

# ------------------------------------------------------- attribute coverage
GROUP_NON_APPLICABLE = {
    "labelImageContains", "labelImageURL", "safetyDataSheet", "numberOfHazardousComponents",
    "hasBatteries", "batterySize", "ionBatteryFormFactor", "ionNumberOfBatteries",
    "ionBatteryModel", "ionIncludedBatteryPackaging", "ionNumberOfBatteryCells",
    "batteryWattHour", "ionBatteryWeight", "metalBatteryWeight", "metalNumberOfBatteries",
    "metalBatteryFormFactor", "metalBatteryModel", "metalIncludedBatteryPackaging",
    "metalBatteryCellCount", "lithiumMetalContentWeight", "required_storage_condition",
    "pesticide_type", "academicInstitution", "prop65WarningText", "character", "character_group",
    "children_product_certificate_document_reference_id",
    "children_product_test_report_document_reference_id", "states", "zipCodes",
    "frameColorConfiguration", "frontEndPhotoPartner",
    "general_certificate_of_conformity_document_reference_id",
    "law_label_registration_number", "law_label_identification_provider",
    "photoPaperFinishConfiguration", "personalRelationship", "photoAccessoryItemSku",
    "photoConfigurationAttributeNames", "photoItemStoreUpc", "photoOrderQuantityTier",
    "sportsLeague", "sportsTeam", "thirdPartyAccreditationSymbolOnProductPackageCode",
    "upperBodyStrapConfiguration", "IsPreorder", "releaseDate", "certification_type",
}

# attributes that are applicable to a women's apparel item but currently empty
FILLABLE = [
    ("Clothing Shoulder Style", "clothing_shoulder_style", "肩部造型", "Button Shoulder / Button Detail", "本款唯一差异化细节（肩部装饰扣）却没落到属性里，属白白浪费的结构化卖点，直接影响肩部风格筛选"),
    ("Collar Style", "collarType", "领型", "Crew Neck / Round Neck", "标题与描述都写了 crew neck，但 collarType 空 → 领型筛选拿不到流量"),
    ("Sleeve Style", "sleeveStyle", "袖型", "Set-in Sleeve（待核对闭集）", "长袖已填 sleeveLengthStyle，缺袖型，袖型筛选缺失"),
    ("Clothing Back Style", "clothing_back_style", "背部造型", "Plain / Pullover Back（待核对闭集）", "补齐服装结构属性，提升完整度权重"),
    ("Manufacturer Name", "manufacturer", "制造商", "实际生产商法定名称", "合规与售后必需；空值会降低 listing 完整度"),
    ("Manufacturer Part Number", "manufacturerPartNumber", "制造商料号", "PNK366085", "可用于款式归一与售后溯源"),
    ("Model Number", "modelNumber", "型号", "PNK366085", "同上；空值影响部分类目筛选"),
    ("Product Line", "productLine", "产品线", "PINKSIR Fall Tunic Collection", "系列化运营，便于站内系列页/关联推荐"),
    ("Theme", "theme", "主题", "Fall / Business Casual（待核对闭集）", "场景主题词可进筛选与推荐池"),
    ("Net Content Statement", "netContentStatement", "净含量声明", "1 Blouse", "包装声明，合规完整性"),
    ("Retail Packaging", "ib_retail_packaging", "零售包装", "Poly Bag", "影响 WFS 入库与包装描述"),
    ("Ships in Original Packaging", "shipsInOriginalPackaging", "原包装发货", "No", "影响配送方式与退货率预期"),
    ("Designer", "designer", "设计师", "PINKSIR（可选）", "品牌资产属性，可留空但填了更完整"),
    ("MSRP", "msrp", "建议零售价", "29.99", "价格锚；$19.99 需要原价对比才显得是“折扣价”"),
    ("Minimum Seller Allowed Price", "minimumSellerAllowedPrice", "最低售价", "18.99", "价格守卫，防止促销互踩"),
    ("Maximum Seller Allowed Price", "maximumSellerAllowedPrice", "最高售价", "24.99", "价格守卫"),
    ("Repricer Strategy", "repricerStrategy", "调价策略", "No Repricing", "不启用自动调价时显式声明，避免被批量比价"),
    ("Recycled Material", "recycledMaterial", "再生材质", "有再生涤纶则填 Yes", "竞品用 OEKO-TEX + Climate Pledge 做环保背书，可再生材质是对标位"),
    ("Percentage of Recycled Material", "percentageOfRecycledMaterial", "再生比例", "按实际（如 30）", "配合上一项"),
    ("Warranty Text", "warrantyText", "质保说明", "如 30 天质保则填", "提升转化与退货率表现"),
    ("Warranty URL", "warrantyURL", "质保链接", "对应说明页链接", "配合上一项"),
    ("SKU Update", "SkuUpdate", "SKU 更新标记", "No", "误填 Yes 会触发 SKU 变更，务必保持 No"),
    ("External Product ID Type", "externalProductIdType", "外部商品 ID 类型", "ASIN（如做跨平台关联）", "跨平台关联用；无则留空"),
    ("External Product ID", "externalProductId", "外部商品 ID", "对应 ASIN", "配合上一项"),
]

total_attr = sum(1 for i, a in enumerate(ids) if i > 2 and (a or labels[i]))
filled_attr = sum(1 for i, a in enumerate(ids)
                  if i > 2 and (a or labels[i]) and s(first[i]) != "")
na_attr = sum(1 for a in ids[3:] if a in GROUP_NON_APPLICABLE)
applicable = total_attr - na_attr
applicable_filled = filled_attr
applicable_empty = applicable - applicable_filled

# ------------------------------------------------------------ keyword data
rev = load(REV_X, "US-B0F93MLVNZ(2000)")
kw_header, kw_rows = rev[0], rev[1:]
words = load(REV_X, "Unique Words")
wf = [(s(r[0]), int(r[1])) for r in words[1:] if s(r[0]) and r[1]]

COMP_TITLE = ("AUTOMET Womens Long Sleeve Shirts Fall Sweaters Business Casual Tops Blouse "
              "Work Tunic Tops Trendy Cute Sweaters Dressy Casual Blouse Plus Size 2026 "
              "Teacher Outfits Fashion Winter Clothes").lower()

STOP = {"the", "and", "for", "with", "women", "womens", "woman", "top", "tops", "women's"}


def covered(text, kw):
    toks = [t for t in re.findall(r"[a-z0-9']+", kw.lower()) if len(t) > 2 and t not in STOP]
    if not toks:
        return False
    return all(t in text for t in toks)


def num(v):
    try:
        return float(str(v).replace(",", "").replace("$", ""))
    except Exception:
        return None


keywords = []
for r in kw_rows:
    if not s(r[0]):
        continue
    keywords.append({
        "w": s(r[0]),
        "cn": s(r[1]),
        "vol": num(r[17]),
        "density": num(r[19]),
        "products": num(r[24]),
        "ratio": num(r[25]),
        "ppc": s(r[30]),
        "rank": s(r[10]),
        "own": covered(OWN_TEXT, s(r[0])),
        "comp": covered(COMP_TITLE, s(r[0])),
        "amz": ("amazon" in s(r[0]).lower() or "storefront" in s(r[0]).lower()
                or "returns" in s(r[0]).lower()),
    })

# competitor compare set (from US-Alexa Products sheet)
alx = load(ALX_X, "Products(6)")
comp_set = []
for r in alx[1:]:
    if len(r) > 8 and s(r[3]):
        comp_set.append({
            "asin": s(r[3]),
            "title": s(r[4]),
            "price": num(r[7]),
            "reviews": s(r[9]),
            "rating": num(r[10]),
        })
alx_qa = load(ALX_X, "Alexa(5)")
qa = [(s(r[1]), s(r[2])) for r in alx_qa[1:] if len(r) > 2 and s(r[0]) == "原文"]

# title candidates
TITLES = [
    "PINKSIR Women's Plus Size Long Sleeve Ribbed Tunic Blouse, Fall Business Casual Top, 1X-5X",
    "PINKSIR Womens Plus Size Ribbed Knit Tunic Blouse, Long Sleeve Fall Work Top, Casual Pullover, 1X-5X",
    "PINKSIR Women's Plus Size Ribbed Tunic Blouse, Long Sleeve Fall Business Casual Work Top, 1X-5X",
]
NEW_BULLETS = [
    "[Fabric] 95% Polyester / 5% Spandex ribbed knit - lightweight stretch with a soft, non-bulky "
    "handfeel that resists wrinkles and keeps its shape all day.",
    "[Plus Size Fit] Made for women's plus sizes 1X, 2X, 3X, 4X and 5X with a relaxed tunic length "
    "that covers the hip and skims the midsection for a flattering line.",
    "[Signature Detail] Decorative shoulder buttons on a classic crew neck add a polished accent "
    "that sets this plus size tunic apart from plain long sleeve tops.",
    "[Occasion] Easy to style for business casual, office, teacher outfits, travel and weekend wear "
    "- pairs with leggings, jeans, trousers or a skirt.",
    "[Layer Friendly] Long sleeves and a smooth knit finish layer cleanly under cardigans, blazers "
    "and vests for fall, winter and transitional spring weather.",
    "[Care & Sizing] Machine wash cold on gentle cycle, do not bleach, hang or lay flat to dry. "
    "Please compare the PINKSIR size chart with a similar top before ordering.",
]

payload = {
    "attr": {"filled": applicable_filled, "empty": applicable_empty, "na": na_attr,
             "total": total_attr},
    "title": {"text": TITLE, "len": len(TITLE)},
    "titles": [{"text": t, "len": len(t)} for t in TITLES],
    "bullets": BULLETS,
    "newBullets": NEW_BULLETS,
    "colors": COLORS, "sizes": SIZES, "price": PRICE,
    "keywords": keywords,
    "wordFreq": wf[:30],
    "compSet": comp_set,
    "ownPrice": float(PRICE),
    "qa": qa,
    "fillable": [{"label": a, "id": b, "cn": c, "val": d, "why": e}
                 for a, b, c, d, e in FILLABLE],
}

echarts_src = open(ECHARTS, "r", encoding="utf-8").read()
TEMPLATE = open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             "listing_report_template.html"), "r", encoding="utf-8").read()
html = (TEMPLATE
        .replace("/*__ECHARTS__*/", echarts_src)
        .replace("__PAYLOAD__", json.dumps(payload, ensure_ascii=False)))
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    f.write(html)
print("WROTE", OUT, len(html), "chars")
print("attr:", payload["attr"], "| title len", len(TITLE))
for t in TITLES:
    print("  title cand", len(t), t)
