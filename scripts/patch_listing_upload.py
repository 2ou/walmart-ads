"""Patch the Walmart MP_WFS_ITEM upload template with the optimized PNK366085 listing.

Writes:
  outputs/PNK366085-Blouses-Tops-upload-optimized-20260916.xlsx   (upload-ready)
  D:\\WorkBuddyData\\walmart-ads\\listing_dump\\patch_report.txt    (change log)
"""
import os
import shutil

from openpyxl import load_workbook

SRC = r"C:\Users\Administrator\Desktop\PNK366085-Blouses-Tops-upload.xlsx"
DST_DIR = r"D:\Tools\Code\walmart-ads\outputs"
DST = os.path.join(DST_DIR, "PNK366085-Blouses-Tops-upload-optimized-20260916.xlsx")
BACKUP_DIR = r"D:\WorkBuddyData\walmart-ads\backup"
LOG = r"D:\WorkBuddyData\walmart-ads\listing_dump\patch_report.txt"

NEW_TITLE = ("PINKSIR Womens Plus Size Tunic Tops Crewneck Long Sleeve Ribbed Blouse "
             "Fall Business Casual, 1X-5X")

NEW_KF = [
    "\U0001F9F6 [Soft Ribbed Knit] - Made from 95% polyester and 5% spandex, the lightweight "
    "ribbed-knit fabric offers comfortable stretch without a heavy sweater feel.",

    "\u2728 [Button Detail] - Decorative shoulder buttons add a polished accent to the simple "
    "solid-color design and set this tunic apart from plain long sleeve tops.",

    "\U0001F45A [Easy Pullover] - The classic crew neck and pull-on construction make this long "
    "sleeve blouse easy to wear, easy to layer and comfortable all day.",

    "\U0001F4CF [Plus Size Fit 1X-5X] - Designed in women's plus sizes 1X, 2X, 3X, 4X and 5X with "
    "a relaxed body and tunic-style silhouette that covers the hip and skims the midsection.",

    "\U0001F4BC [Business Casual & Work] - A dressy casual fall top for the office, teacher "
    "outfits, meetings and travel. Pair this plus size tunic top with jeans, leggings, work "
    "pants or a skirt.",

    "\U0001F9FC [Simple Care] - Machine wash cold on a gentle cycle or hand wash. Do not bleach; "
    "hang or lay flat to dry. Please compare the PINKSIR size chart with a similar top before "
    "ordering.",
]

NEW_DESC = (
    "<p><strong>Comfortable style for workdays and weekends.</strong> This PINKSIR women's plus "
    "size tunic blouse combines a relaxed silhouette with subtle details that make everyday "
    "outfits feel more polished. The lightweight ribbed-knit fabric is made from 95% polyester "
    "and 5% spandex, providing comfortable stretch without the weight or bulk of a traditional "
    "sweater.</p>\n\n"
    "<p>Decorative shoulder buttons add a refined accent to the classic crew neck, while the "
    "pull-on construction makes this fall top easy to wear and layer. Long sleeves provide "
    "comfortable coverage for fall, winter and transitional spring weather. The tunic-style "
    "silhouette works well worn loose over leggings or jeans and can also be paired with "
    "trousers for a clean business-casual look that fits an office, teacher or travel "
    "wardrobe.</p>\n\n"
    "<ul>\n"
    "  <li>Lightweight ribbed-knit texture with comfortable stretch</li>\n"
    "  <li>Classic crew neck and long sleeves</li>\n"
    "  <li>Decorative shoulder button detail</li>\n"
    "  <li>Relaxed plus-size fit available in sizes 1X-5X</li>\n"
    "  <li>Suitable for fall, business casual, work, travel and everyday layering</li>\n"
    "</ul>\n\n"
    "<p>For care, machine wash cold on a gentle cycle or hand wash. Do not bleach. Hang or lay "
    "flat to dry. Please review the PINKSIR size chart and compare the garment measurements "
    "with a similar top before ordering.</p>"
)

SIMPLE = {
    "productName": NEW_TITLE,
    "shortDescription": NEW_DESC,
    "manufacturer": "PINKSIR",
    "manufacturerPartNumber": "PNK366085",
    "modelNumber": "PNK366085",
    "productLine": "PINKSIR Fall Tunic Collection",
    "theme": "Fall",
    "netContentStatement": "1 Blouse",
    "designer": "PINKSIR",
    "msrp": "29.99",
    "minimumSellerAllowedPrice": "18.99",
    "maximumSellerAllowedPrice": "24.99",
    "shipsInOriginalPackaging": "No",
    "SkuUpdate": "No",
}

SKIP_CLOSED_SET = {
    "collarType": "Crew Neck 已由 shirtNeckStyle 承载，无需重复",
    "sleeveStyle": "闭集值需在卖家中心下拉核对（本款为常规长袖）",
    "clothing_back_style": "闭集值需核对（本款为常规后背）",
    "clothing_shoulder_style": "肩扣已由 embellishmentType=Ornamental Button 承载，闭集无 Button Shoulder",
    "ib_retail_packaging": "闭集值需核对",
    "repricerStrategy": "预置列表，需在卖家中心选择",
    "recycledMaterial": "无再生材质数据，留空",
    "percentageOfRecycledMaterial": "无再生材质数据，留空",
    "warrantyText": "has_written_warranty=No，无质保政策，留空",
    "warrantyURL": "同上，留空",
    "externalProductIdType": "无跨平台关联，留空",
    "externalProductId": "无跨平台关联，留空",
}

L = []

os.makedirs(BACKUP_DIR, exist_ok=True)
os.makedirs(DST_DIR, exist_ok=True)
shutil.copy2(SRC, os.path.join(BACKUP_DIR, os.path.basename(SRC)))
L.append("backup -> %s" % os.path.join(BACKUP_DIR, os.path.basename(SRC)))

wb = load_workbook(SRC)
ws = wb["Product Content And Site Exp"]

# --- locate structure -------------------------------------------------------
id_row = None
for r in range(1, min(ws.max_row, 30) + 1):
    if str(ws.cell(row=r, column=4).value or "").strip() == "sku":
        id_row = r
        break
assert id_row, "attribute-id row not found"
data_start = id_row + 2
data_end = ws.max_row

ids = {}
for c in range(1, ws.max_column + 1):
    a = ws.cell(row=id_row, column=c).value
    if a:
        ids.setdefault(str(a).strip(), []).append(c)

L.append("attribute-id row=%d | data rows %d..%d | cols=%d" % (id_row, data_start, data_end, ws.max_column))
L.append("")

changes = {}


def put(attr, value, n=1, note=""):
    cols = ids.get(attr)
    if not cols:
        L.append("  !! %s NOT FOUND - skipped" % attr)
        return
    c = cols[n - 1]
    for r in range(data_start, data_end + 1):
        if str(ws.cell(row=r, column=4).value or "").strip():
            ws.cell(row=r, column=c).value = value
    changes.setdefault(attr, {"col": c, "n": n, "value": value, "note": note})


def put_multi(attr, values, note=""):
    cols = ids.get(attr)
    if not cols:
        L.append("  !! %s NOT FOUND - skipped" % attr)
        return
    for i, v in enumerate(values):
        c = cols[i]
        for r in range(data_start, data_end + 1):
            if str(ws.cell(row=r, column=4).value or "").strip():
                ws.cell(row=r, column=c).value = v
    changes.setdefault(attr, {"col": cols[0], "n": len(values), "value": values, "note": note})


# --- 1. title ---------------------------------------------------------------
put("productName", NEW_TITLE, note="71 -> %d chars" % len(NEW_TITLE))
# --- 2. key features --------------------------------------------------------
put_multi("keyFeatures", NEW_KF, note="6 bullets rewritten")
# --- 3. site description ----------------------------------------------------
put("shortDescription", NEW_DESC, note="keyword-enriched, structure kept")
# --- 4. variant second axis color+clothingSizeGroup -> color+clothingSize ----
cols = ids["variantAttributeNames"]
assert len(cols) >= 2
for r in range(data_start, data_end + 1):
    if str(ws.cell(row=r, column=4).value or "").strip():
        ws.cell(row=r, column=cols[1]).value = "clothingSize"
changes["variantAttributeNames[2]"] = {"col": cols[1], "n": 1,
                                       "value": "clothingSize",
                                       "note": "clothingSizeGroup(Woman's Plus constant) -> clothingSize (1X-5X)"}
# --- 5. attribute fill ------------------------------------------------------
for k, v in SIMPLE.items():
    if k in ("productName", "shortDescription"):
        continue
    put(k, v)

# --- report -----------------------------------------------------------------
L.append("=== CHANGES APPLIED (all %d data rows) ===" % (data_end - data_start + 1))
for k, d in changes.items():
    if isinstance(d["value"], list):
        L.append("  %-34s col=%-4s | %s" % (k, d["col"], d["note"]))
        for i, v in enumerate(d["value"]):
            L.append("      [%d] len=%-4d %s" % (i + 1, len(v), v[:90]))
    else:
        v = d["value"]
        L.append("  %-34s col=%-4s | len=%-4d | %s | %s" % (
            k, d["col"], len(str(v)), str(v)[:70], d["note"]))

L.append("")
L.append("=== LEFT EMPTY (needs manual pick from closed list / real data) ===")
for k, why in SKIP_CLOSED_SET.items():
    col = ids.get(k, ["?"])[0]
    L.append("  %-34s col=%-4s | %s" % (k, col, why))

L.append("")
L.append("=== VERIFY: read back first data row ===")
for k in list(SIMPLE.keys()) + ["variantAttributeNames"]:
    if k not in ids:
        continue
    c = ids[k][0]
    L.append("  %-34s = %s" % (k, str(ws.cell(row=data_start, column=c).value)[:80]))
c2 = ids["variantAttributeNames"][1]
L.append("  %-34s = %s" % ("variantAttributeNames[1]", ws.cell(row=data_start, column=c2).value))

# --- title length guard -----------------------------------------------------
L.append("")
L.append("title length = %d (limit 100, ideal <90)" % len(NEW_TITLE))
for i, b in enumerate(NEW_KF):
    L.append("  bullet %d len=%d" % (i + 1, len(b)))

wb.save(DST)
L.append("")
L.append("SAVED -> %s" % DST)

os.makedirs(os.path.dirname(LOG), exist_ok=True)
with open(LOG, "w", encoding="utf-8") as f:
    f.write("\n".join(L))
print("DONE", DST)
