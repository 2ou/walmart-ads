"""Dump full text of the fields we plan to rewrite, plus the other visible sheet's header."""
import os
from openpyxl import load_workbook

SRC = r"C:\Users\Administrator\Desktop\PNK366085-Blouses-Tops-upload.xlsx"
OUT = r"D:\WorkBuddyData\walmart-ads\listing_dump\own_text_full.txt"


def s(v):
    return "" if v is None else str(v).strip()


wb = load_workbook(SRC, read_only=True)
rows = [list(r) for r in wb["Product Content And Site Exp"].iter_rows(values_only=True)]
other = None
try:
    other = [list(r) for r in wb["Trade Item Configurations"].iter_rows(values_only=True, max_row=6)]
except Exception as e:
    other = [["ERR", str(e)]]
wb.close()

id_idx = next(i for i, r in enumerate(rows) if len(r) > 3 and s(r[3]) == "sku")
ids = [s(x) for x in rows[id_idx]]
labels = [s(x) for x in rows[id_idx - 1]]
first = rows[id_idx + 2]

L = []


def find_col(name, n=1):
    c = 0
    for i, a in enumerate(ids):
        if a == name:
            c += 1
            if c == n:
                return i
    return None


def show(attr, n=1, title=None):
    ci = find_col(attr, n)
    if ci is None:
        L.append("!! %s #%d NOT FOUND" % (attr, n))
        return
    L.append("--- %s (%s) col=%d len=%d ---" % (title or labels[ci], attr, ci, len(s(first[ci]))))
    L.append(s(first[ci]))
    L.append("")


show("productName", 1, "TITLE")
show("shortDescription", 1, "SITE DESCRIPTION")
for k in range(1, 7):
    show("keyFeatures", k, "KEY FEATURE %d" % k)

L.append("=== all isPrimaryVariant values by row ===")
pc = find_col("isPrimaryVariant")
skc = find_col("sku")
for r in range(id_idx + 2, len(rows)):
    if s(rows[r][skc]):
        L.append("  %-38s %s" % (s(rows[r][skc]), s(rows[r][pc])))

L.append("")
L.append("=== 'Trade Item Configurations' sheet first rows ===")
for i, r in enumerate(other or []):
    L.append("  row%d: %s" % (i + 1, [s(x)[:28] for x in r[:14]]))

L.append("")
L.append("=== data validation count ===")
L.append("  n/a (read_only)")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(L))
print("WROTE", OUT)
