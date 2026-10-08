#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
导入沃尔玛 WFS 仓储费 / 配送费明细到 walmart_ads 库。

数据源目录：D:\\PNK\\06WOSZ财务\\<YYYY.M> 沃尔玛PNK财务数据详情\\
  - 6.仓储费用明细表\\*.xlsx          -> wm_ads_wfs_storage_fee
  - 7.WFS配送费报表\\*.xlsx           -> wm_ads_wfs_fulfillment_fee  (2026.01 ~ 2026.03)
  - 7.配送费用明细表\\*.xlsx          -> wm_ads_wfs_fulfillment_fee  (2026.04 起改名)

设计要点：
  1. 两类报表前几行是元信息（公司名 / 说明 / 周期汇总），明细表头位置逐月固定但需探测：
       storage    -> 表头行找第一列 == "GTIN"，其后 1 行是字段说明（跳过），再往后才是数据
       settlement -> 表头行找第一列以 "Walmart.com Order" 开头，下一行即数据
  2. 列名逐月有差异（1 月无 Billing Method；6/7 月多 Original Amount / Discount Savings；
     7 月起长期仓储费拆 366-450 / 451+ 两档，8 月列名又微调），一律按列名规则映射，不按下标。
  3. 日期为 Excel 序列号（46018 -> 2025-12-27），统一换算。
  4. settlement 只有 GTIN 没有 SKU，用同期 storage 文件的 GTIN->SKU 字典反补。
  5. 幂等：全部走 INSERT ... ON DUPLICATE KEY UPDATE，重复执行不会产生脏数据。

用法：
  python import_wfs_fees.py                    # 导入 >= 2026-01 的全部月份
  python import_wfs_fees.py --dry-run          # 只解析统计，不写库
  python import_wfs_fees.py --start 2026-06    # 只导入 2026-06 及以后
  python import_wfs_fees.py --months 2026-07,2026-08
"""

import os
import re
import sys
import glob
import codecs
import argparse
import datetime

import openpyxl
import pymysql

# ---------------------------------------------------------------- 配置

DEFAULT_ROOT = r"D:\PNK\06WOSZ财务"
DEFAULT_START = "2026-01"

DB_CFG = dict(host="127.0.0.1", port=3306, user="root", password="root",
              database="walmart_ads", charset="utf8mb4")

# 月份文件夹：如 "2026.01 沃尔玛PNK财务数据详情" / "2026.4 沃尔玛PNK财务数据详情"
DIR_RE = re.compile(r"^(\d{4})\.(\d{1,2})\s+.*沃尔玛PNK财务数据详情$")

EXCEL_EPOCH = datetime.date(1899, 12, 30)

STORAGE_DIR = "6.仓储费用明细表"
FULFILL_DIRS = ("7.WFS配送费报表", "7.配送费用明细表")

# V38 建表脚本（与 Flyway 共用同一份，脚本用 IF NOT EXISTS，重复执行无害）
DDL_PATH = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..",
    "src", "main", "resources", "db", "migration", "V38__wfs_fee_tables.sql"))

# ---------------------------------------------------------------- 工具


def xldate(v):
    """Excel 序列号 -> date；失败返回 None。"""
    if v is None or v == "":
        return None
    try:
        return EXCEL_EPOCH + datetime.timedelta(days=int(float(v)))
    except (TypeError, ValueError):
        return None


def dec(v):
    """转 Decimal 友好的 float；失败返回 None。"""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("$", "")
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def txt(v, limit=None):
    if v is None:
        return None
    s = str(v).strip()
    if s == "":
        return None
    return s[:limit] if limit else s


def parent_of(sku):
    """PNK366069-C-Blue-XL-WFS -> PNK366069"""
    if not sku:
        return None
    return str(sku).split("-")[0].strip() or None


def norm(name):
    return re.sub(r"\s+", " ", str(name or "")).strip().lower()


# ---------------------------------------------------------------- 列名映射

STORAGE_FIXED = {
    "gtin": "gtin",
    "sku": "sku",
    "item id": "item_id",
    "item name": "item_name",
    "length (in)": "length_in",
    "width (in)": "width_in",
    "height (in)": "height_in",
    "weight (lb)": "weight_lb",
    "days in report period": "days_in_period",
    "final storage fee": "final_storage_fee",
    "original amount": "original_amount",
    "discount savings": "discount_savings",
}

LT2_RE = re.compile(r"45[01]\s*\+")  # 451+ / 450+  -> 超长期档


def map_storage_col(raw):
    n = norm(raw)
    if n in STORAGE_FIXED:
        return STORAGE_FIXED[n]
    if "daily storage fee per unit" in n:
        if n.startswith("standard"):
            return "std_daily_fee"
        if n.startswith("peak"):
            return "peak_daily_fee"
        return "lt2_daily_fee" if LT2_RE.search(n) else "lt_daily_fee"
    if "average units available" in n or "average daily units available" in n:
        if n.startswith("standard"):
            return "std_avg_units"
        if n.startswith("peak"):
            return "peak_avg_units"
        return "lt2_avg_units" if LT2_RE.search(n) else "lt_avg_units"
    return None  # 如 "Long-term (366-450): total units" 不落库


FULFILL_MAP = {
    "walmart.com order #": "order_id",
    "walmart.com order line #": "order_line",
    "walmart.com po #": "po_no",
    "walmart.com po line #": "po_line",
    "transaction type": "transaction_type",
    "reason code": "reason_code",
    "detail": "detail",
    "wfsreferenceid": "wfs_reference_id",
    "transaction date/time": "txn_date",
    "qty": "qty",
    "billing method": "billing_method",
    "partner id": "partner_id",
    "partner gtin": "partner_gtin",
    "partner item name": "partner_item_name",
    "category": "category",
    "sub category": "sub_category",
    "net payable": "net_payable",
    "original amount": "original_amount",
    "discount savings": "discount_savings",
    "customerorderid": "customer_order_id",
    "orderchannelid": "order_channel_id",
}


def map_fulfill_col(raw):
    return FULFILL_MAP.get(norm(raw))


# ---------------------------------------------------------------- 文件解析


def find_header(rows, pred, max_scan=30):
    for i, r in enumerate(rows[:max_scan]):
        if r and r[0] is not None and pred(str(r[0]).strip()):
            return i
    return None


def parse_storage(path):
    """返回 (period_start, period_end, total_fee_declared, [row dict, ...])"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    hi = find_header(rows, lambda s: s.upper() == "GTIN")
    if hi is None:
        raise RuntimeError("未找到 GTIN 表头: %s" % path)

    # 定位 "Report start date / Report end date ..." 汇总头，其下一行是汇总值
    period_start = period_end = None
    declared_total = None
    for i, r in enumerate(rows[:hi]):
        if r and r[0] is not None and norm(r[0]) == "report start date":
            if i + 1 < len(rows):
                s = rows[i + 1]
                period_start = xldate(s[0])
                period_end = xldate(s[1])
                declared_total = dec(s[2]) if len(s) > 2 else None
            break

    hdr = rows[hi]
    colmap = {}
    for idx, raw in enumerate(hdr):
        key = map_storage_col(raw)
        if key and key not in colmap:
            colmap[key] = idx

    out = []
    for r in rows[hi + 2:]:                       # 跳过 hi+1 的字段说明行
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        rec = {k: (r[i] if i < len(r) else None) for k, i in colmap.items()}
        sku = txt(rec.get("sku"), 180)
        rec["sku"] = sku
        rec["parent_sku"] = parent_of(sku)
        rec["item_name"] = txt(rec.get("item_name"), 600)
        rec["gtin"] = txt(rec.get("gtin"), 32)
        out.append(rec)
    return period_start, period_end, declared_total, out


def parse_fulfillment(path):
    """返回 (period_start, period_end, [row dict, ...])"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    hi = find_header(rows, lambda s: s.startswith("Walmart.com Order"))
    if hi is None:
        raise RuntimeError("未找到 Order 表头: %s" % path)

    # 定位 "ReportingPeriod_StartDate / EndDate" 行，其下一行是周期值
    period_start = period_end = None
    for i, r in enumerate(rows[:hi]):
        if r and r[0] is not None and norm(r[0]).startswith("reportingperiod_startdate"):
            if i + 1 < len(rows):
                period_start = xldate(rows[i + 1][0])
                period_end = xldate(rows[i + 1][1])
            break

    hdr = rows[hi]
    colmap = {}
    for idx, raw in enumerate(hdr):
        key = map_fulfill_col(raw)
        if key and key not in colmap:
            colmap[key] = idx

    out = []
    for r in rows[hi + 1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        rec = {k: (r[i] if i < len(r) else None) for k, i in colmap.items()}
        out.append(rec)
    return period_start, period_end, out


# ---------------------------------------------------------------- 扫描

def scan_months(root, start_label=None, only=None):
    """返回 [(period_label, storage_files, fulfill_files), ...]"""
    items = []
    for name in sorted(os.listdir(root)):
        d = os.path.join(root, name)
        if not os.path.isdir(d):
            continue
        m = DIR_RE.match(name)
        if not m:
            continue
        label = "%s-%02d" % (m.group(1), int(m.group(2)))
        if start_label and label < start_label:
            continue
        if only and label not in only:
            continue
        s_files = sorted(glob.glob(os.path.join(d, STORAGE_DIR, "*.xlsx")))
        f_files = []
        for sub in FULFILL_DIRS:
            f_files += sorted(glob.glob(os.path.join(d, sub, "*.xlsx")))
        s_files = [f for f in s_files if "~$" not in os.path.basename(f)]
        f_files = [f for f in f_files if "~$" not in os.path.basename(f)]
        items.append((label, s_files, f_files))
    return items


# ---------------------------------------------------------------- 写库

STORAGE_COLS = ["period_start", "period_end", "period_label", "gtin", "sku", "parent_sku",
                "item_id", "item_name", "length_in", "width_in", "height_in", "weight_lb",
                "std_daily_fee", "peak_daily_fee", "lt_daily_fee", "lt2_daily_fee",
                "std_avg_units", "peak_avg_units", "lt_avg_units", "lt2_avg_units",
                "days_in_period", "final_storage_fee", "original_amount", "discount_savings",
                "source_file"]

FULFILL_COLS = ["period_start", "period_end", "period_label", "order_id", "order_line",
                "po_no", "po_line", "transaction_type", "reason_code", "detail",
                "wfs_reference_id", "txn_date", "qty", "billing_method", "partner_id",
                "partner_gtin", "sku", "parent_sku", "partner_item_name", "category",
                "sub_category", "net_payable", "original_amount", "discount_savings",
                "customer_order_id", "order_channel_id", "source_file"]

DEC_FIELDS = {"length_in", "width_in", "height_in", "weight_lb",
              "std_daily_fee", "peak_daily_fee", "lt_daily_fee", "lt2_daily_fee",
              "std_avg_units", "peak_avg_units", "lt_avg_units", "lt2_avg_units",
              "final_storage_fee", "original_amount", "discount_savings",
              "qty", "net_payable"}
INT_FIELDS = {"days_in_period"}
DATE_FIELDS = {"txn_date"}
TXT_LIMITS = {"item_name": 600, "partner_item_name": 600, "sku": 180,
              "wfs_reference_id": 300, "reason_code": 200, "detail": 200,
              "source_file": 300}


def build_row(rec, cols, ctx):
    row = []
    for c in cols:
        if c in ("period_start", "period_end", "period_label", "source_file"):
            row.append(ctx[c])
            continue
        v = rec.get(c)
        if c in DATE_FIELDS:
            v = xldate(v)
        elif c in DEC_FIELDS:
            v = dec(v)
        elif c in INT_FIELDS:
            d = dec(v)
            v = int(d) if d is not None else None
        else:
            v = txt(v, TXT_LIMITS.get(c))
            if c == "parent_sku" and not v:
                v = parent_of(rec.get("sku"))
        row.append(v)
    return row


def split_sql_statements(text):
    """按分号切分 SQL，但忽略字符串字面量与行注释中的分号。"""
    stmts, buf, in_q = [], [], False
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if in_q:
            buf.append(ch)
            if ch == "\\" and i + 1 < n:
                buf.append(text[i + 1])
                i += 2
                continue
            if ch == "'":
                in_q = False
            i += 1
            continue
        if ch == "'":
            in_q = True
            buf.append(ch)
        elif ch == "#" or text.startswith("--", i):
            while i < n and text[i] != "\n":
                i += 1
            continue
        elif ch == ";":
            s = "".join(buf).strip()
            if s:
                stmts.append(s)
            buf = []
        else:
            buf.append(ch)
        i += 1
    s = "".join(buf).strip()
    if s:
        stmts.append(s)
    return stmts


def init_ddl(conn):
    """执行 V38 建表脚本（幂等）。用于尚未启动过应用、Flyway 未跑到的环境。"""
    if not os.path.exists(DDL_PATH):
        raise RuntimeError("找不到建表脚本: %s" % DDL_PATH)
    text = codecs.open(DDL_PATH, encoding="utf-8").read()
    count = 0
    with conn.cursor() as cur:
        for stmt in split_sql_statements(text):
            cur.execute(stmt)
            count += 1
    conn.commit()
    print("已执行建表脚本 V38（%d 条语句）" % count)


def upsert_many(conn, table, cols, rows, update_cols):
    if not rows:
        return 0
    ph = ", ".join(["%s"] * len(cols))
    sql = "INSERT INTO `%s` (%s) VALUES (%s) ON DUPLICATE KEY UPDATE %s" % (
        table,
        ", ".join("`%s`" % c for c in cols),
        ph,
        ", ".join("`%s`=VALUES(`%s`)" % (c, c) for c in update_cols),
    )
    with conn.cursor() as cur:
        cur.executemany(sql, rows)
    conn.commit()
    return len(rows)


# ---------------------------------------------------------------- 主流程


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=DEFAULT_ROOT)
    ap.add_argument("--start", default=DEFAULT_START, help="起始月份，如 2026-01")
    ap.add_argument("--months", default="", help="只处理指定月份，逗号分隔，如 2026-07,2026-08")
    ap.add_argument("--init-ddl", action="store_true",
                    help="导入前先执行 V38 建表脚本（表不存在时必选）")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    only = set(x.strip() for x in args.months.split(",") if x.strip()) or None
    months = scan_months(args.root, args.start, only)
    if not months:
        print("未匹配到任何月份文件夹：%s" % args.root)
        return 1

    print("待处理月份：%s" % ", ".join(m[0] for m in months))

    # ---- 第一遍：解析全部仓储费，顺便建 GTIN -> SKU 字典
    gtin_sku = {}
    storage_batches = []
    for label, s_files, _ in months:
        if not s_files:
            print("  [WARN] %s 缺仓储费文件" % label)
            continue
        for f in s_files:
            ps, pe, declared, recs = parse_storage(f)
            if ps is None and recs:
                raise RuntimeError("周期解析失败: %s" % f)
            for r in recs:
                if r.get("gtin") and r.get("sku"):
                    gtin_sku[r["gtin"]] = r["sku"]
            total = sum(dec(r.get("final_storage_fee")) or 0 for r in recs)
            storage_batches.append((label, f, ps, pe, declared, total, recs))
            flag = "OK " if declared is None or abs((declared or 0) - total) < 0.05 else "!! "
            print("  %s仓储费 %s | %s~%s | %d 行 | 明细合计 %.2f | 表头 Total %s"
                  % (flag, label, ps, pe, len(recs), total, declared))

    # ---- 第二遍：解析配送费
    fulfill_batches = []
    for label, _, f_files in months:
        if not f_files:
            print("  [WARN] %s 缺配送费文件" % label)
            continue
        for f in f_files:
            ps, pe, recs = parse_fulfillment(f)
            ff = sum(dec(r.get("net_payable")) or 0 for r in recs
                     if txt(r.get("transaction_type")) == "FulfillmentFee")
            fulfill_batches.append((label, f, ps, pe, ff, recs))
            print("   OK 配送费 %s | %s~%s | %d 行 | FulfillmentFee 合计 %.2f"
                  % (label, ps, pe, len(recs), ff))

    print("\nGTIN->SKU 字典：%d 条" % len(gtin_sku))

    unmapped = {}
    for _, _, _, _, _, recs in fulfill_batches:
        for r in recs:
            g = txt(r.get("partner_gtin"), 32)
            if g and g not in gtin_sku:
                unmapped[g] = unmapped.get(g, 0) + 1
    if unmapped:
        print("配送费中未匹配到 SKU 的 GTIN：%d 个（sku 置空，不影响金额）" % len(unmapped))
        for g, c in sorted(unmapped.items(), key=lambda x: -x[1])[:10]:
            print("    %s x%d" % (g, c))

    if args.dry_run:
        print("\n[dry-run] 未写库。")
        return 0

    # ---- 写库
    conn = pymysql.connect(**DB_CFG)
    try:
        if args.init_ddl:
            init_ddl(conn)
        n_s = n_f = 0
        for label, f, ps, pe, declared, total, recs in storage_batches:
            ctx = dict(period_start=ps, period_end=pe, period_label=label,
                       source_file=os.path.basename(f))
            rows = [build_row(r, STORAGE_COLS, ctx) for r in recs]
            n_s += upsert_many(conn, "wm_ads_wfs_storage_fee", STORAGE_COLS, rows,
                               [c for c in STORAGE_COLS if c not in ("period_start", "sku")])

        for label, f, ps, pe, ff, recs in fulfill_batches:
            ctx = dict(period_start=ps, period_end=pe, period_label=label,
                       source_file=os.path.basename(f))
            rows = []
            for r in recs:
                g = txt(r.get("partner_gtin"), 32)
                sku = gtin_sku.get(g)
                r = dict(r)
                r["sku"] = sku
                r["parent_sku"] = parent_of(sku)
                rows.append(build_row(r, FULFILL_COLS, ctx))
            n_f += upsert_many(conn, "wm_ads_wfs_fulfillment_fee", FULFILL_COLS, rows,
                               [c for c in FULFILL_COLS if c != "wfs_reference_id"])

        print("\n写入完成：仓储费 %d 行，配送费 %d 行" % (n_s, n_f))
    finally:
        conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
