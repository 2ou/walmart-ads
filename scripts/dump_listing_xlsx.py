"""Dump the 3 referenced workbooks to text so the listing analysis can be done.

Usage: python dump_listing_xlsx.py
Writes D:/WorkBuddyData/walmart-ads/listing_dump/<name>.txt  (one per sheet)
"""
import os
import sys
from openpyxl import load_workbook

OUT = r"D:\WorkBuddyData\walmart-ads\listing_dump"
os.makedirs(OUT, exist_ok=True)

FILES = [
    r"C:\Users\Administrator\Desktop\PNK366085-Blouses-Tops-upload.xlsx",
    r"C:\Users\Administrator\Desktop\ReverseASIN-US-B0F93MLVNZ(2000)-20260916.xlsx",
    r"C:\Users\Administrator\Desktop\US-Alexa(5)-20260916.xlsx",
]


def cell(v):
    if v is None:
        return ""
    s = str(v)
    s = s.replace("\r\n", " ⏎ ").replace("\n", " ⏎ ").replace("\t", " ")
    return s.strip()


def main():
    for path in FILES:
        if not os.path.exists(path):
            print("MISSING", path)
            continue
        wb = load_workbook(path, data_only=True, read_only=True)
        base = os.path.splitext(os.path.basename(path))[0]
        print("=" * 100)
        print("FILE:", base)
        print("SHEETS:", wb.sheetnames)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            # trim fully-empty trailing rows
            while rows and all(cell(c) == "" for c in rows[-1]):
                rows.pop()
            safe = "".join(ch if ch.isalnum() or ch in "-_()" else "_" for ch in ws.title)
            fn = os.path.join(OUT, f"{base}__{safe}.txt")
            with open(fn, "w", encoding="utf-8") as f:
                f.write(f"FILE: {base}\nSHEET: {ws.title}\nROWS: {len(rows)}  COLS: {ws.max_column}\n")
                f.write("-" * 80 + "\n")
                for i, r in enumerate(rows, 1):
                    vals = [cell(c) for c in r]
                    while vals and vals[-1] == "":
                        vals.pop()
                    if not vals:
                        continue
                    f.write(f"[{i}] " + " | ".join(vals) + "\n")
            print(f"  wrote {fn}  rows={len(rows)} cols={ws.max_column}")
        wb.close()


if __name__ == "__main__":
    main()
